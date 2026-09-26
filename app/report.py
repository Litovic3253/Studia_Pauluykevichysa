"""Смета за период (месяц, квартал или свои даты) — Excel-файл с заказами, деньгами и прибылью.

Лист «Смета»: каждый заказ периода (по дате создания, отменённые не входят) — цена, предоплата,
получено, долг, себестоимость (пластик по закупке), прибыль; внизу строка «Итого»
формулами Excel. Лист «Итоги»: сводка по деньгам и расход пластика по материалам."""
import calendar
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import db
import pricing

MONTHS_NOM = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь",
              "Ноябрь", "Декабрь"]

COLUMNS = [
    # (заголовок, ширина, формат чисел, суммировать в «Итого»)
    ("№", 6, None, False),
    ("Дата", 11, None, False),
    ("Клиент", 22, None, False),
    ("Описание", 30, None, False),
    ("Материал", 10, None, False),
    ("Пластик, г", 11, "#,##0.0", True),
    ("Печать, ч", 10, "#,##0.00", True),
    ("Кол-во", 8, "0", True),
    ("Статус", 14, None, False),
    ("Цена", 12, "#,##0", True),
    ("Предоплата", 12, "#,##0", True),
    ("Оплачен", 9, None, False),
    ("Получено", 12, "#,##0", True),
    ("Долг", 12, "#,##0", True),
    ("Себестоимость (пластик)", 15, "#,##0", True),
    ("Прибыль", 12, "#,##0", True),
]


# ---------- периоды ----------

def month_period(year: int, month: int) -> tuple[date, date]:
    return date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1])


def quarter_period(year: int, quarter: int) -> tuple[date, date]:
    first = 3 * (quarter - 1) + 1
    return date(year, first, 1), month_period(year, first + 2)[1]


def period_title(start: date, end: date) -> str:
    """«Сентябрь 2026», «3 квартал 2026» или «01.09.2026 – 15.09.2026»."""
    if start.day == 1 and (start, end) == month_period(start.year, start.month):
        return f"{MONTHS_NOM[start.month - 1]} {start.year}"
    if start.day == 1 and start.month % 3 == 1 and (start, end) == quarter_period(start.year, start.month // 3 + 1):
        return f"{start.month // 3 + 1} квартал {start.year}"
    return f"{start.strftime('%d.%m.%Y')} – {end.strftime('%d.%m.%Y')}"


def default_file_name(start: date, end: date) -> str:
    title = period_title(start, end).replace(" – ", "_").replace(" ", "_")
    return f"Смета_{title}.xlsx"


# ---------- данные ----------

def report_rows(start: date, end: date) -> list[dict]:
    settings = db.get_settings()
    rows = []
    for o in db.orders_between(start.isoformat(), end.isoformat()):
        qty = o["qty"] or 1
        exp = db.expense(o["material"], o["weight_g"] or 0, o["print_hours"] or 0, qty, settings)
        price = o["price"] or 0
        prepayment = o["prepayment"] or 0
        received = price if o["paid"] else prepayment
        rows.append({
            "id": o["id"], "date": o["created_at"][:10], "client": o["client"] or "",
            "description": o["description"] or "", "material": o["material"] or "",
            "grams": (o["weight_g"] or 0) * qty, "hours": (o["print_hours"] or 0) * qty, "qty": qty,
            "status": db.STATUSES.get(o["status"], o["status"]).split(" ", 1)[-1],
            "price": price, "prepayment": prepayment, "paid": bool(o["paid"]),
            "received": received, "debt": max(price - received, 0),
            "cost": exp["total"],
            "profit": price - exp["total"],
        })
    return rows


def summary(rows: list[dict]) -> dict:
    total = {k: sum(r[k] for r in rows) for k in
             ("price", "received", "debt", "cost", "profit", "grams", "hours")}
    total["orders"] = len(rows)
    total["margin"] = total["profit"] / total["price"] if total["price"] else 0
    by_material: dict[str, float] = {}
    for r in rows:
        by_material[r["material"] or "—"] = by_material.get(r["material"] or "—", 0) + r["grams"]
    total["by_material"] = sorted(by_material.items(), key=lambda kv: -kv[1])
    return total


# ---------- Excel ----------

HEAD_FILL = PatternFill("solid", fgColor="18181B")
TOTAL_FILL = PatternFill("solid", fgColor="F4F4F5")
THIN = Side(style="thin", color="E4E4E7")


def _put(ws, row: int, col: int, value):
    """Пишет значение; текст пользователя, начинающийся с «=», остаётся текстом, а не формулой."""
    cell = ws.cell(row, col, value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def write_report(path: str | Path, start: date, end: date) -> Path:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_name(path.name + ".xlsx")
    rows = report_rows(start, end)
    total = summary(rows)
    title = period_title(start, end)
    wb = Workbook()

    ws = wb.active
    ws.title = "Смета"
    ws["A1"] = f"Смета · {title}"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = (f"Период {start.strftime('%d.%m.%Y')} – {end.strftime('%d.%m.%Y')} · заказы по дате создания, "
                f"без отменённых · валюта: {db.get_settings()['currency']}")
    ws["A2"].font = Font(size=9, color="71717A")
    head_row = 4
    for col, (name, width, _, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(head_row, col, name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[head_row].height = 30

    keys = ["id", "date", "client", "description", "material", "grams", "hours", "qty", "status", "price",
            "prepayment", "paid", "received", "debt", "cost", "profit"]
    for i, r in enumerate(rows):
        row = head_row + 1 + i
        for col, key in enumerate(keys, start=1):
            value = r[key]
            if key == "date":
                value = pricing.fmt_date(value)
            elif key == "paid":
                value = "да" if value else "нет"
            elif isinstance(value, float):
                value = round(value, 2)
            cell = _put(ws, row, col, value)
            fmt = COLUMNS[col - 1][2]
            if fmt:
                cell.number_format = fmt
            cell.border = Border(bottom=THIN)

    first, last = head_row + 1, head_row + len(rows)
    total_row = last + 1 if rows else head_row + 1
    ws.cell(total_row, 1, "Итого").font = Font(bold=True)
    for col, (_, _, fmt, summable) in enumerate(COLUMNS, start=1):
        cell = ws.cell(total_row, col)
        cell.fill = TOTAL_FILL
        if summable and rows:
            letter = get_column_letter(col)
            cell.value = f"=SUM({letter}{first}:{letter}{last})"
            cell.number_format = fmt
            cell.font = Font(bold=True)
    ws.freeze_panes = ws.cell(head_row + 1, 4)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(COLUMNS))}{max(last, head_row)}"

    s = wb.create_sheet("Итоги")
    s.column_dimensions["A"].width = 34
    s.column_dimensions["B"].width = 18
    s["A1"] = f"Итоги · {title}"
    s["A1"].font = Font(size=14, bold=True)
    lines = [
        ("Заказов", total["orders"], "0"),
        ("Сумма заказов", total["price"], "#,##0"),
        ("Получено (оплаты + предоплаты)", total["received"], "#,##0"),
        ("Ждём доплату", total["debt"], "#,##0"),
        ("Себестоимость (пластик по закупке)", total["cost"], "#,##0"),
        ("Прибыль", total["profit"], "#,##0"),
        ("Маржа", total["margin"], "0%"),
        ("Пластика израсходовано, кг", total["grams"] / 1000, "0.00"),
        ("Печать, часов", total["hours"], "0.0"),
    ]
    for i, (label, value, fmt) in enumerate(lines, start=3):
        s.cell(i, 1, label)
        cell = s.cell(i, 2, round(value, 4))
        cell.number_format = fmt
        if label == "Прибыль":
            s.cell(i, 1).font = cell.font = Font(bold=True)
    row = len(lines) + 5
    s.cell(row, 1, "Расход пластика по материалам").font = Font(bold=True)
    for i, (name, grams) in enumerate(total["by_material"], start=row + 1):
        _put(s, i, 1, name)
        s.cell(i, 2, round(grams / 1000, 3)).number_format = '0.000 "кг"'

    wb.save(path)
    return path
