"""Запись заказов/клиентов/цен в .xlsx через openpyxl."""
from pathlib import Path

from openpyxl import Workbook

from app.export import build_all_sheets


def export_to_excel(path: str | Path) -> Path:
    """Сохраняет все листы в .xlsx и возвращает итоговый путь (с расширением .xlsx)."""
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_name(path.name + ".xlsx")
    wb = Workbook()
    wb.remove(wb.active)  # убираем дефолтный пустой лист "Sheet"
    for sheet_name, rows in build_all_sheets().items():
        ws = wb.create_sheet(title=sheet_name)
        for row in rows:
            ws.append(row)
            # openpyxl считает строку, начинающуюся с "=", формулой — пользовательский
            # текст (имя клиента, заметки) должен остаться текстом.
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
    wb.save(path)
    return path
