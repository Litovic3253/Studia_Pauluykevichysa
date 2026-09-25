"""Тесты для app/excel_writer.py — запись .xlsx через openpyxl."""
import openpyxl

import db
from app import excel_writer


def test_export_to_excel_creates_file_with_three_sheets(temp_db, tmp_path):
    db.add_order({"client": "Иван", "contact": "", "material": "PLA",
                   "weight_g": 100, "print_hours": 1, "qty": 1, "cost": 100, "price": 200})
    dest = tmp_path / "export.xlsx"

    excel_writer.export_to_excel(dest)

    assert dest.exists()
    wb = openpyxl.load_workbook(dest)
    assert set(wb.sheetnames) == {"Заказы", "Клиенты", "Цены"}


def test_export_to_excel_writes_order_row_correctly(temp_db, tmp_path):
    order_id = db.add_order({"client": "Мария", "contact": "@maria", "material": "PETG",
                              "weight_g": 50, "print_hours": 2, "qty": 1, "cost": 300, "price": 500})
    dest = tmp_path / "export.xlsx"

    excel_writer.export_to_excel(dest)

    wb = openpyxl.load_workbook(dest)
    ws = wb["Заказы"]
    header = [cell.value for cell in ws[1]]
    data_row = [cell.value for cell in ws[2]]
    assert header[0] == "ID"
    assert data_row[header.index("ID")] == order_id
    assert data_row[header.index("Клиент")] == "Мария"
