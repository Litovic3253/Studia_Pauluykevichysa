"""Смета за период, автоточки в дате, поле с календарём, боковое меню."""
from datetime import date, datetime

import flet as ft
from openpyxl import load_workbook

import db
import pricing
from app import input_hints, report
from app.sidebar import Action, Destination, Sidebar


def test_format_date_input_adds_dots_while_typing():
    assert pricing.format_date_input("27092026") == "27.09.2026"
    assert pricing.format_date_input("2709") == "27.09"
    assert pricing.format_date_input("270") == "27.0"
    assert pricing.format_date_input("27") == "27"
    assert pricing.format_date_input("270920261") == "27.09.2026"  # лишняя цифра отбрасывается


def test_format_date_input_leaves_words_and_own_formats_alone():
    assert pricing.format_date_input("завтра") == "завтра"
    assert pricing.format_date_input("1.1.26") == "1.1.26"
    assert pricing.format_date_input("27.09.2026") == "27.09.2026"
    assert pricing.format_date_input("") == ""


def test_date_field_formats_on_change_and_fills_from_calendar():
    picked = []
    field = input_hints.date_field("Срок", on_pick=lambda: picked.append(field.value))
    field.value = "01102026"
    field.on_change(None)
    assert field.value == "01.10.2026"

    picker = field.data
    assert isinstance(picker, ft.DatePicker)
    picker.value = datetime(2026, 12, 31)
    picker.on_change(None)
    assert field.value == "31.12.2026"
    assert picked == ["31.12.2026"]


def test_periods_and_titles():
    assert report.month_period(2026, 2) == (date(2026, 2, 1), date(2026, 2, 28))
    assert report.quarter_period(2026, 3) == (date(2026, 7, 1), date(2026, 9, 30))
    assert report.period_title(*report.month_period(2026, 9)) == "Сентябрь 2026"
    assert report.period_title(*report.quarter_period(2026, 4)) == "4 квартал 2026"
    assert report.period_title(date(2026, 9, 3), date(2026, 9, 10)) == "03.09.2026 – 10.09.2026"


def _order(created: str, **fields) -> int:
    oid = db.add_order({"client": "Иван", "contact": "", "material": "PLA", "weight_g": 100, "print_hours": 2,
                         "qty": 1, "cost": 0, "price": 1000, **fields})
    with db._conn() as c:
        c.execute("UPDATE orders SET created_at = ? WHERE id = ?", (created + "T12:00:00", oid))
    return oid


def test_report_includes_only_period_orders_and_computes_profit(temp_db, tmp_path):
    db.set_setting("purchase_prices", {"PLA": 1500})
    db.set_setting("hour_rate", 50)
    paid = _order("2026-09-05")
    db.update_order(paid, paid=1)
    _order("2026-09-20", prepayment=300)
    cancelled = _order("2026-09-21")
    db.update_order(cancelled, status="cancelled")
    _order("2026-10-01")  # другой месяц

    start, end = report.month_period(2026, 9)
    rows = report.report_rows(start, end)
    assert [r["id"] for r in rows] == [paid, paid + 1]
    total = report.summary(rows)
    assert total["price"] == 2000
    assert total["received"] == 1000 + 300
    assert total["debt"] == 700
    assert total["cost"] == 2 * 150
    assert total["profit"] == 2000 - 300

    path = report.write_report(tmp_path / "смета", start, end)
    assert path.suffix == ".xlsx"
    wb = load_workbook(path)
    assert wb.sheetnames == ["Смета", "Итоги"]
    ws = wb["Смета"]
    assert ws["A1"].value == "Смета · Сентябрь 2026"
    assert ws.cell(7, 1).value == "Итого"
    assert str(ws.cell(7, 10).value).startswith("=SUM(J5:J6")


def test_report_keeps_formula_like_text_as_text(temp_db, tmp_path):
    _order("2026-09-05", description="=HYPERLINK(1)")
    path = report.write_report(tmp_path / "r.xlsx", *report.month_period(2026, 9))
    assert load_workbook(path)["Смета"].cell(5, 4).value == "=HYPERLINK(1)"


def test_sidebar_selects_item_and_reports_change():
    changes = []
    bar = Sidebar([Destination(ft.Icons.HOME, ft.Icons.HOME, "Один"),
                   Destination(ft.Icons.HOME, ft.Icons.HOME, "Два")], on_change=changes.append)
    clicked = []
    bar.set_layout(True, None, [Action(ft.Icons.SAVE, "Сохранить", clicked.append)])
    items = bar.content.controls
    items[1].on_click("click")
    assert bar.selected_index == 1 and changes == ["click"]
    bar.content.controls[-1].on_click("x")
    assert clicked == ["x"]
    bar.set_layout(False, None, [])
    assert bar.width == 72 and bar.content.controls[0].tooltip == "Один"
